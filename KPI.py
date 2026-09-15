import tkinter as tk
from tkinter import ttk, filedialog
from datetime import datetime, timedelta
import sys
import ctypes
from calendar import monthrange
import gspread
from oauth2client.service_account import ServiceAccountCredentials
import os
import time
import threading
import tempfile
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Font, Alignment
from openpyxl.utils import get_column_letter

try:
    from reconciliation_queue import enqueue_job, is_reconciliation_task
except ImportError:
    enqueue_job = None
    is_reconciliation_task = lambda t: False

try:
    from kpi_profile import export_filename, load_saved_profile
except ImportError:
    export_filename = lambda u, s, e: f"KPI_{u}_{s:%Y-%m-%d}_to_{e:%Y-%m-%d}.xlsx"
    load_saved_profile = lambda *a, **k: None

# Windows dark title bar support
if sys.platform == "win32":
    try:
        ctypes.windll.dwmapi.DwmSetWindowAttribute(
            ctypes.windll.user32.GetParent(0), 2, ctypes.byref(ctypes.c_int(2)), 4
        )
    except Exception:
        pass

# ==========================================
# CONSTANTS & THEME (Aurora Borealis & Obsidian)
# ==========================================
BASE_BG = '#faf8f2'
CARD_BG = '#ffffff'
CARD_BORDER = '#e2e8f0'
CARD_BORDER_HOVER = '#00e5ff'

ACCENT_DARK = '#0b1420'
ACCENT_HOVER = '#162e4c'
HEADER_BG = '#0b1420'
HEADER_BTN_BG = '#162e4c'
HEADER_BTN_HOVER = '#1f3f66'

AURORA_CYAN = '#00e5ff'
AURORA_MINT = '#49cf9e'
AURORA_AMBER = '#ff6b00'
AURORA_SAPPHIRE = '#00b7ff'

TEXT_MAIN = '#0f172a'
TEXT_MUTED = '#64748b'
TEXT_LIGHT = '#94a3b8'
TEXT_WHITE = '#ffffff'

INPUT_BG = '#ffffff'
INPUT_BORDER = '#cbd5e1'
INPUT_BORDER_FOCUS = '#00b7ff'
PLACEHOLDER_COLOR = '#94a3b8'

BTN_PRIMARY_BG = '#0d1b2a'
BTN_PRIMARY_FG = '#ffffff'
BTN_PRIMARY_HOVER = '#162e4c'

BTN_SUBMIT_BG = '#10b981'
BTN_SUBMIT_HOVER = '#059669'
BTN_DELETE_BG = '#ef4444'
BTN_DELETE_HOVER = '#dc2626'
BTN_SECONDARY_BG = '#f1f5f9'
BTN_SECONDARY_HOVER = '#e2e8f0'

STATUS_PILL_ONLINE_BG = '#e2fdf2'
STATUS_PILL_ONLINE_BORDER = '#49cf9e'
STATUS_PILL_ONLINE_FG = '#065f46'

STATUS_PILL_OFFLINE_BG = '#fef2f2'
STATUS_PILL_OFFLINE_BORDER = '#ef4444'
STATUS_PILL_OFFLINE_FG = '#991b1b'

# Compatibility aliases
BG_DARK = BASE_BG
BG_PANEL = CARD_BG
BG_ENTRY = INPUT_BG
ORANGE = AURORA_AMBER
GREEN = BTN_SUBMIT_BG
RED = BTN_DELETE_BG
BLUE = '#0284c7'
TEXT_GRAY = TEXT_MUTED

MSG_SUCCESS = '#059669'
MSG_ERROR = '#dc2626'
MSG_INFO = '#0284c7'
MSG_WARNING = '#d97706'

FONT_LABEL = ('Segoe UI', 9, 'bold')
FONT_ENTRY = ('Segoe UI', 10)
FONT_BUTTON = ('Segoe UI', 9, 'bold')
FONT_HEADER = ('Segoe UI', 15, 'bold')
FONT_COUNT = ('Segoe UI', 9, 'bold')
FONT_MSG = ('Segoe UI', 9)

# ============================================
# USER ID & CODE DEFAULTS
# ============================================
USER_ID = "Tharindu"
USER_CODE = "CSSUN156"
# ============================================

SPREADSHEET_URL = "https://docs.google.com/spreadsheets/d/1FyX5TQgoVluPYfFol6L0M813Aqov8eoC6fRk4lp1hn0/edit?usp=sharing"
SHEET1_NAME = "Sheet1"
SHEET2_NAME = "Sheet2"

CREDENTIALS_FILES = ["credentials.json", "credentials", "credentials.txt"]

# ============================================
# EXCEL EXPORT CONFIG
# ============================================
EXCEL_FILENAME = "UTILIZATION UPDATE 2026 - Copy.xlsx"
EXCEL_FOLDER = ""
TARGET_SHEET = "RECONCILIATION"

# Row layout
META_HEADER_ROW  = 6
TASK_HEADER_ROW  = 7
FIRST_DATA_ROW   = 8

# Column layout
DATE_COL, USER_COL, INTIME_COL, UCODE_COL, SITE_COL = 1, 2, 3, 4, 5
TASK_COL_START = 6
TASK_COL_STEP  = 3

DAYTIME_SLOTS = ["8:00 - 12:00", "12:01 - 15:00", "15:01 - 17:00"]
OT_SLOT = "OT"
WEIGHTS = [0.50, 0.30, 0.20]

TASK_TO_HEADER = {
    "GDN Reconciliation:":    ["GDN Reconciliation:"],
    "GRN Reconciliation:":    ["GRN Reconciliation:"],
    "GDN Creation:":          ["GDN Creation:"],
    "GRN Creation:":          ["GRN Creation:"],
    "Load Plan or Asn:":      ["Load Plan or Asn:"],
    "Load Audit:":            ["Load Audit:"],
    "Shipping:":              ["Shipping:"],
    "Load Transfer:":         ["Load Transfer:", "Load Transfer:"],
    "Allocation/Backorders:": ["Allocation", "Backorders"],
}

YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
BLUE_FILL   = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
THIN_SIDE   = Side(style="thin", color="000000")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def excel_path():
    if EXCEL_FOLDER and EXCEL_FOLDER.strip():
        folder = EXCEL_FOLDER.strip()
    else:
        folder = os.path.dirname(os.path.abspath(__file__))
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, EXCEL_FILENAME)


def ensure_excel_exists(path):
    """If the Excel file doesn't exist, create it with the correct header layout."""
    if os.path.exists(path):
        return True
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = TARGET_SHEET

        ws.cell(row=2, column=5).value = "TOTAL"
        ws.cell(row=3, column=5).value = "CLIENT"
        ws.cell(row=4, column=5).value = "SYSTEM"
        ws.cell(row=5, column=5).value = "UOM"
        ws.cell(row=6, column=5).value = "T-CODE"
        ws.cell(row=7, column=5).value = "SITE"

        ws.cell(row=7, column=1).value = "DATE"
        ws.cell(row=7, column=2).value = "USER"
        ws.cell(row=7, column=3).value = "IN-TIME"
        ws.cell(row=7, column=4).value = "U-CODE"

        tasks = [
            ("GRN Reconciliation:", 6,  "Inbound",  "Per Reconciliation Report", "CSSTR0324"),
            ("GRN Creation:",        9,  "Inbound",  "Per vehicle",               "CSSTR0325"),
            ("GDN Reconciliation:",  12, "Outbound", "Per Reconciliation Report", "CSSTR0324"),
            ("GDN Creation:",        15, "Outbound", "Per vehicle",               "CSSTR0326"),
            ("Load Plan or Asn:",    18, "Outbound", "Per load plan",             "CSSTR0327"),
            ("Load Audit:",          21, "Outbound", "Load ID Count",             ""),
            ("Shipping:",            24, "Outbound", "Load ID Count",             ""),
            ("Load Transfer:",       27, "Outbound", "Load ID Count",             ""),
            ("Load Transfer:",       30, "Outbound", "LP Count",                  ""),
            ("Allocation",           33, "Outbound", "LP Count",                  ""),
            ("Backorders",           36, "Outbound", "LP Count",                  ""),
        ]
        for name, col, system, uom, tcode in tasks:
            ws.cell(row=4, column=col).value = system
            ws.cell(row=5, column=col).value = uom
            if tcode:
                ws.cell(row=6, column=col).value = tcode
            ws.cell(row=7, column=col).value = name

        wb.save(path)
        print(f"[EXCEL] Created new file: {path}")
        return True
    except Exception as e:
        print(f"[EXCEL] Could not create file: {e}")
        return False


def _norm_header(s):
    return str(s).replace(":", "").replace("\n", " ").strip().lower()


def find_header_cols(ws, header_texts):
    wanted = [_norm_header(h) for h in header_texts]
    found = []
    used = set()
    for want in wanted:
        hit = None
        for c in range(TASK_COL_START, ws.max_column + 1, TASK_COL_STEP):
            if c in used:
                continue
            v = ws.cell(row=TASK_HEADER_ROW, column=c).value
            if v is not None and _norm_header(v) == want:
                hit = c
                break
        if hit is None:
            return None
        used.add(hit)
        found.append(hit)
    return found


def find_user_rows(ws, iso_date, user_id):
    rows = {}
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        d = ws.cell(row=r, column=DATE_COL).value
        u = ws.cell(row=r, column=USER_COL).value
        t = ws.cell(row=r, column=INTIME_COL).value
        if d is None or u is None or t is None:
            continue
        if hasattr(d, "date"):
            d_cmp = d.date()
        else:
            d_cmp = str(d).strip()
            if d_cmp:
                for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%m-%Y"):
                    try:
                        d_cmp = datetime.strptime(d_cmp, fmt).date()
                        break
                    except ValueError:
                        continue
        if str(d_cmp) != str(iso_date):
            continue
        if str(u).strip().lower() != user_id.strip().lower():
            continue
        rows[str(t).strip()] = r
    return rows


def create_missing_rows(ws, iso_date, user_id, user_code):
    existing = find_user_rows(ws, iso_date, user_id)
    if existing:
        return False
    next_row = ws.max_row + 1
    while next_row <= ws.max_row + 20 and ws.cell(row=next_row, column=DATE_COL).value not in (None, ""):
        next_row += 1
    d_obj = datetime.strptime(iso_date, "%Y-%m-%d")
    date_cell = datetime(d_obj.year, d_obj.month, d_obj.day).date()
    for i, slot in enumerate(DAYTIME_SLOTS + [OT_SLOT]):
        r = next_row + i
        cell = ws.cell(row=r, column=DATE_COL)
        cell.value = date_cell
        cell.number_format = "m/d/yyyy"
        ws.cell(row=r, column=USER_COL).value = user_id
        ws.cell(row=r, column=INTIME_COL).value = slot
        ws.cell(row=r, column=UCODE_COL).value = user_code
    return True


def split_count(n):
    if n <= 0:
        return [0, 0, 0]
    raw = [n * w for w in WEIGHTS]
    floors = [int(x) for x in raw]
    remainder = n - sum(floors)
    order = sorted(range(3), key=lambda i: raw[i] - floors[i], reverse=True)
    parts = floors[:]
    for i in range(remainder):
        parts[order[i]] += 1
    return parts


def _apply_border_and_ot_highlight(ws, ot_row_indices, last_col):
    last_row = ws.max_row
    for r in range(1, last_row + 1):
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).border = THIN_BORDER

    for r in range(1, 8):
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).fill = BLUE_FILL

    for r in ot_row_indices:
        for c in range(1, last_col + 1):
            ws.cell(row=r, column=c).fill = YELLOW_FILL


def _find_last_task_col(ws):
    last = 5
    for c in range(TASK_COL_START, ws.max_column + 1, TASK_COL_STEP):
        v = ws.cell(row=TASK_HEADER_ROW, column=c).value
        if v is not None and str(v).strip() != "":
            last = c
    return last + (TASK_COL_STEP - 1)


def write_counts_to_excel(path, user_id, user_code, counts_by_day):
    if not os.path.exists(path):
        return False, f"Excel file not found:\n{path}", []
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        return False, f"Could not open Excel: {e}", []
    if TARGET_SHEET not in wb.sheetnames:
        return False, f"Sheet '{TARGET_SHEET}' not found.", []
    ws = wb[TARGET_SHEET]

    details = []
    any_written = False
    ot_rows_to_highlight = []

    for iso_date, day_counts in sorted(counts_by_day.items()):
        create_missing_rows(ws, iso_date, user_id, user_code)
        rows = find_user_rows(ws, iso_date, user_id)
        if not rows:
            details.append(f"⚠ No rows for {user_id} on {iso_date}")
            continue

        if OT_SLOT in rows:
            ot_rows_to_highlight.append(rows[OT_SLOT])

        for label, val in day_counts.items():
            if label not in TASK_TO_HEADER:
                continue
            cols = find_header_cols(ws, TASK_TO_HEADER[label])
            if not cols:
                details.append(f"⚠ {label} — header not found")
                continue
            if isinstance(val, tuple):
                load_val, lp_val = val
            else:
                load_val, lp_val = val, None
            parts_load = split_count(int(load_val)) if load_val else [0, 0, 0]
            parts_lp   = split_count(int(lp_val)) if lp_val is not None else None
            for i, slot in enumerate(DAYTIME_SLOTS):
                r = rows.get(slot)
                if r is None:
                    continue
                ws.cell(row=r, column=cols[0]).value = parts_load[i]
                if parts_lp is not None and len(cols) > 1:
                    ws.cell(row=r, column=cols[1]).value = parts_lp[i]
                any_written = True

        details.append(f"✔ {iso_date} written")

    last_col = _find_last_task_col(ws)
    _apply_border_and_ot_highlight(ws, ot_rows_to_highlight, last_col)
    details.append(f"✔ Borders applied (A1:{get_column_letter(last_col)}{ws.max_row})")
    details.append(f"✔ {len(ot_rows_to_highlight)} OT row(s) highlighted yellow")
    details.append(f"✔ Blue header (rows 1-7) applied")

    try:
        ws.sheet_view.zoomScale = 80
        ws.sheet_view.zoomScaleNormal = 80
    except Exception as e:
        print(f"[ZOOM] Could not set zoom: {e}")

    try:
        wb.save(path)
    except Exception as e:
        return False, f"Could not save Excel: {e}", details
    return True, "Excel updated.", details


def build_counts_for_range(values1, values2, user_id, start_date, end_date,
                            count_sheet1_fn, sum_sheet2_fn):
    result = {}
    d = start_date
    while d <= end_date:
        date_str = d.strftime("%d-%m-%Y")
        iso = d.strftime("%Y-%m-%d")
        day = {}

        for label in TASK_TO_HEADER.keys():
            if label in ("Load Transfer:", "Allocation/Backorders:", "Load Audit:", "Shipping:"):
                continue
            if label == "GDN Creation:":
                day[label] = (count_sheet1_fn(values1, "GDN Creation:", user_id, date_str)
                              + count_sheet1_fn(values1, "GDN Reconciliation:", user_id, date_str))
            elif label == "GRN Creation:":
                day[label] = (count_sheet1_fn(values1, "GRN Creation:", user_id, date_str)
                              + count_sheet1_fn(values1, "GRN Reconciliation:", user_id, date_str))
            else:
                day[label] = count_sheet1_fn(values1, label, user_id, date_str)

        load, _ = sum_sheet2_fn(values2, "Load Audit:", user_id, date_str)
        day["Load Audit:"] = load

        load, _ = sum_sheet2_fn(values2, "Shipping:", user_id, date_str)
        day["Shipping:"] = load

        l1, l2 = sum_sheet2_fn(values2, "Load Transfer:", user_id, date_str)
        day["Load Transfer:"] = (l1, l2)

        a1, a2 = sum_sheet2_fn(values2, "Allocation/Backorders:", user_id, date_str)
        day["Allocation/Backorders:"] = (a1, a2)

        result[iso] = day
        d += timedelta(days=1)
    return result


def compute_export_range(selected_date):
    if selected_date.day == 1:
        first_of_this_month = selected_date.replace(day=1)
        last_of_prev        = first_of_this_month - timedelta(days=1)
        first_of_prev       = last_of_prev.replace(day=1)
        return first_of_prev, last_of_prev
    else:
        start = selected_date.replace(day=1)
        end   = selected_date - timedelta(days=1)
        return start, end


def generate_user_export(folder, user_id, user_code, selected_date, values1, values2,
                         count_sheet1_fn, sum_sheet2_fn):
    start_date, end_date = compute_export_range(selected_date)
    path = os.path.join(folder, export_filename(user_id, start_date, end_date))
    fd, temp_path = tempfile.mkstemp(prefix='.kpi_', suffix='.xlsx', dir=folder)
    os.close(fd)
    os.unlink(temp_path)
    try:
        if not ensure_excel_exists(temp_path):
            return False, 'Could not create Excel workbook', path, []
        counts = build_counts_for_range(values1, values2, user_id, start_date, end_date,
                                        count_sheet1_fn, sum_sheet2_fn)
        ok, message, details = write_counts_to_excel(temp_path, user_id, user_code, counts)
        if not ok:
            return False, message, path, details
        os.replace(temp_path, path)
        return True, message, path, details
    except Exception as e:
        return False, str(e), path, []
    finally:
        if os.path.exists(temp_path):
            try:
                os.unlink(temp_path)
            except Exception:
                pass


class GoogleSheetsManager:
    def __init__(self, spreadsheet_url):
        self.spreadsheet_url = spreadsheet_url
        self.client = None
        self.sheet = None
        self.worksheet1 = None
        self.worksheet2 = None
        self.connected = False
        self.cred_file_used = None
        self.last_request_time = 0
        self.min_request_interval = 0.1

    def _rate_limit_wait(self):
        current_time = time.time()
        time_since_last = current_time - self.last_request_time
        if time_since_last < self.min_request_interval:
            time.sleep(self.min_request_interval - time_since_last)
        self.last_request_time = time.time()

    def connect(self):
        try:
            script_dir = os.path.dirname(os.path.abspath(__file__))
            for cred_file in CREDENTIALS_FILES:
                file_path = os.path.join(script_dir, cred_file)
                if os.path.exists(file_path):
                    self.cred_file_used = file_path
                    break
                else:
                    cwd_path = os.path.join(os.getcwd(), cred_file)
                    if os.path.exists(cwd_path):
                        self.cred_file_used = cwd_path
                        break

            if not self.cred_file_used:
                print("Credentials file not found (looked for: credentials.json)")
                self.connected = False
                return False

            scope = ["https://spreadsheets.google.com/feeds",
                    "https://www.googleapis.com/auth/drive"]
            creds = ServiceAccountCredentials.from_json_keyfile_name(self.cred_file_used, scope)
            self.client = gspread.authorize(creds)
            self.sheet = self.client.open_by_url(self.spreadsheet_url)
            self.worksheet1 = self.sheet.worksheet(SHEET1_NAME)
            self.worksheet2 = self.sheet.worksheet(SHEET2_NAME)
            self.connected = True
            print("Connected to Google Sheets!")
            return True
        except Exception as e:
            print(f"Failed to connect: {type(e).__name__} - {e}")
            self.connected = False
            return False

    def get_all_values_sheet1(self):
        if not self.connected or not self.worksheet1:
            return []
        try:
            return self.worksheet1.get_all_values()
        except Exception as e:
            print(f"get_all_values_sheet1 error: {e}")
            return []

    def get_all_values_sheet2(self):
        if not self.connected or not self.worksheet2:
            return []
        try:
            return self.worksheet2.get_all_values()
        except Exception as e:
            print(f"get_all_values_sheet2 error: {e}")
            return []

    def append_row_sheet1(self, row_data):
        if not self.connected or not self.worksheet1:
            return False
        try:
            self._rate_limit_wait()
            self.worksheet1.append_row(row_data)
            return True
        except Exception as e:
            print(f"append_row_sheet1 error: {e}")
            return False

    def append_row_sheet2(self, row_data):
        if not self.connected or not self.worksheet2:
            return False
        try:
            self._rate_limit_wait()
            self.worksheet2.append_row(row_data)
            return True
        except Exception as e:
            print(f"append_row_sheet2 error: {e}")
            return False

    def search_and_delete_row_sheet1(self, search_column, search_value, user_id=None):
        if not self.connected or not self.worksheet1:
            return False
        try:
            self._rate_limit_wait()
            all_values = self.worksheet1.get_all_values()
            for i, row in enumerate(all_values, 1):
                if len(row) > search_column and row[search_column] == str(search_value):
                    if user_id is not None:
                        row_user = row[4].strip() if len(row) > 4 and row[4] else ""
                        if row_user.lower() != str(user_id).lower():
                            continue
                    self.worksheet1.delete_rows(i)
                    return True
            return False
        except Exception as e:
            print(f"search_and_delete_row_sheet1 error: {e}")
            return False

    def search_and_delete_row_sheet2(self, search_column, search_value, user_id=None):
        if not self.connected or not self.worksheet2:
            return False
        try:
            self._rate_limit_wait()
            all_values = self.worksheet2.get_all_values()
            for i, row in enumerate(all_values, 1):
                if len(row) > search_column and row[search_column] == str(search_value):
                    if user_id is not None:
                        row_user = row[5].strip() if len(row) > 5 and row[5] else ""
                        if row_user.lower() != str(user_id).lower():
                            continue
                    self.worksheet2.delete_rows(i)
                    return True
            return False
        except Exception as e:
            print(f"search_and_delete_row_sheet2 error: {e}")
            return False


# ==========================================
# DATEPICKER WIDGET
# ==========================================
class DatePicker:
    def __init__(self, master=None, callback=None):
        self.master = master
        self.callback = callback
        self.selected_date = datetime.now()
        self.view_date = datetime(self.selected_date.year, self.selected_date.month, 1)
        self.calendar_visible = False
        self.calendar_window = None

        self.frame = tk.Frame(master, bg='#ffffff', highlightbackground=CARD_BORDER, highlightthickness=1, cursor='hand2')

        self.cal_icon = tk.Label(
            self.frame, text='📅', bg='#ffffff', fg=TEXT_MUTED, font=('Segoe UI', 9), cursor='hand2'
        )
        self.cal_icon.pack(side='left', padx=(6, 2), pady=3)

        self.date_label = tk.Label(
            self.frame, text=self.selected_date.strftime('%d-%m-%Y'),
            bg='#ffffff', fg=TEXT_MAIN, font=('Segoe UI', 9, 'bold'), cursor='hand2'
        )
        self.date_label.pack(side='left', padx=(2, 4), pady=3)

        self.dropdown_btn = tk.Label(
            self.frame, text='▾', bg='#ffffff', fg=TEXT_LIGHT, font=('Segoe UI', 8), cursor='hand2'
        )
        self.dropdown_btn.pack(side='left', padx=(0, 6), pady=3)

        for w in (self.frame, self.cal_icon, self.date_label, self.dropdown_btn):
            w.bind('<Button-1>', self.toggle_calendar)
            w.bind('<Enter>', lambda e: self._on_enter())
            w.bind('<Leave>', lambda e: self._on_leave())

    def _on_enter(self):
        for w in (self.frame, self.cal_icon, self.date_label, self.dropdown_btn):
            w.config(bg='#f8fafc')

    def _on_leave(self):
        for w in (self.frame, self.cal_icon, self.date_label, self.dropdown_btn):
            w.config(bg='#ffffff')

    def toggle_calendar(self, event=None):
        if self.calendar_visible:
            self.hide_calendar()
        else:
            self.show_calendar()

    def show_calendar(self):
        if self.calendar_window and self.calendar_window.winfo_exists():
            self.hide_calendar()
            return

        self.calendar_window = tk.Toplevel(self.master)
        self.calendar_window.geometry('240x265')
        self.calendar_window.configure(bg=CARD_BG)
        self.calendar_window.resizable(False, False)
        self.calendar_window.overrideredirect(True)

        try:
            x = self.date_label.winfo_rootx() - 10
            y = self.date_label.winfo_rooty() + self.date_label.winfo_height() + 8
            self.calendar_window.geometry(f"+{x}+{y}")
        except Exception:
            pass

        cal_border = tk.Frame(self.calendar_window, bg=CARD_BG, highlightbackground=CARD_BORDER, highlightthickness=1)
        cal_border.pack(fill='both', expand=True)

        header_frame = tk.Frame(cal_border, bg=ACCENT_DARK, padx=6, pady=6)
        header_frame.pack(fill='x')

        prev_btn = tk.Label(header_frame, text='◀', bg=ACCENT_DARK, fg=TEXT_WHITE, font=('Segoe UI', 9, 'bold'), cursor='hand2')
        prev_btn.pack(side='left', padx=4)
        prev_btn.bind('<Button-1>', self.prev_month)

        self.month_label = tk.Label(header_frame, text=self.view_date.strftime('%B %Y'), bg=ACCENT_DARK, fg=TEXT_WHITE, font=('Segoe UI', 9, 'bold'))
        self.month_label.pack(side='left', expand=True)

        next_btn = tk.Label(header_frame, text='▶', bg=ACCENT_DARK, fg=TEXT_WHITE, font=('Segoe UI', 9, 'bold'), cursor='hand2')
        next_btn.pack(side='right', padx=4)
        next_btn.bind('<Button-1>', self.next_month)

        day_frame = tk.Frame(cal_border, bg='#f8fafc', pady=3)
        day_frame.pack(fill='x')
        for d_name in ('Mo', 'Tu', 'We', 'Th', 'Fr', 'Sa', 'Su'):
            lbl = tk.Label(day_frame, text=d_name, bg='#f8fafc', fg=TEXT_MUTED, font=('Segoe UI', 8, 'bold'), width=3)
            lbl.pack(side='left', padx=1)

        self.cal_grid_frame = tk.Frame(cal_border, bg=CARD_BG, padx=4, pady=2)
        self.cal_grid_frame.pack(fill='both', expand=True)

        today_frame = tk.Frame(cal_border, bg='#f8fafc', padx=6, pady=4)
        today_frame.pack(fill='x')
        today_btn = tk.Button(
            today_frame, text='Today', bg='#ffffff', fg=TEXT_MAIN, activebackground='#f1f5f9',
            activeforeground=TEXT_MAIN, relief='flat', highlightbackground=CARD_BORDER, highlightthickness=1,
            font=('Segoe UI', 8, 'bold'), cursor='hand2', command=self.go_to_today
        )
        today_btn.pack(fill='x')

        self.calendar_window.bind('<FocusOut>', lambda e: self.hide_calendar(e))
        self.calendar_window.bind('<Escape>', lambda e: self.hide_calendar(e))

        self.calendar_visible = True
        self.update_calendar()
        self.calendar_window.focus_set()

    def hide_calendar(self, event=None):
        if self.calendar_window and self.calendar_window.winfo_exists():
            self.calendar_window.destroy()
        self.calendar_window = None
        self.calendar_visible = False

    def go_to_today(self):
        now = datetime.now()
        self.view_date = datetime(now.year, now.month, 1)
        self.select_date(now)

    def update_calendar(self):
        for widget in self.cal_grid_frame.winfo_children():
            widget.destroy()

        if hasattr(self, 'month_label') and self.month_label.winfo_exists():
            self.month_label.config(text=self.view_date.strftime('%B %Y'))

        year = self.view_date.year
        month = self.view_date.month
        first_weekday, days_in_month = monthrange(year, month)
        today = datetime.now().date()

        row = 0
        col = 0
        for _ in range(first_weekday):
            lbl = tk.Label(self.cal_grid_frame, text='', bg=CARD_BG, width=3, height=1)
            lbl.grid(row=row, column=col, padx=1, pady=1)
            col += 1

        for day in range(1, days_in_month + 1):
            date_obj = datetime(year, month, day)
            is_today = (date_obj.date() == today)
            is_selected = (date_obj.date() == self.selected_date.date())

            if is_selected:
                bg_color = ACCENT_DARK
                fg_color = TEXT_WHITE
            elif is_today:
                bg_color = '#e0f2fe'
                fg_color = '#0284c7'
            else:
                bg_color = CARD_BG
                fg_color = TEXT_MAIN

            day_btn = tk.Label(
                self.cal_grid_frame, text=str(day), bg=bg_color, fg=fg_color,
                font=('Segoe UI', 9, 'bold') if is_selected else ('Segoe UI', 9),
                width=3, height=1, cursor='hand2', relief='flat'
            )
            day_btn.grid(row=row, column=col, padx=1, pady=1)

            if is_today and not is_selected:
                day_btn.configure(highlightbackground='#00b7ff', highlightthickness=1)

            day_btn.bind('<Button-1>', lambda e, d=date_obj: self.select_date(d))
            if not is_selected:
                day_btn.bind('<Enter>', lambda e, b=day_btn: b.config(bg='#f1f5f9'))
                day_btn.bind('<Leave>', lambda e, b=day_btn, bgc=bg_color: b.config(bg=bgc))

            col += 1
            if col > 6:
                col = 0
                row += 1

    def select_date(self, date_obj):
        self.selected_date = date_obj
        self.date_label.config(text=date_obj.strftime('%d-%m-%Y'))
        if self.calendar_visible and self.calendar_window and self.calendar_window.winfo_exists():
            self.update_calendar()
        if self.callback:
            self.callback(date_obj)
        self.hide_calendar()

    def prev_month(self, event=None):
        year, month = self.view_date.year, self.view_date.month
        self.view_date = datetime(year - 1, 12, 1) if month == 1 else datetime(year, month - 1, 1)
        self.update_calendar()

    def next_month(self, event=None):
        year, month = self.view_date.year, self.view_date.month
        self.view_date = datetime(year + 1, 1, 1) if month == 12 else datetime(year, month + 1, 1)
        self.update_calendar()

    def get_date(self):
        return self.selected_date

    def set_date(self, date_obj):
        self.selected_date = date_obj
        self.date_label.config(text=date_obj.strftime('%d-%m-%Y'))


# ==========================================
# CUSTOM ENTRY & COMBOBOX WIDGETS
# ==========================================
class PlaceholderEntry(tk.Entry):
    def __init__(self, master=None, placeholder="", width=None, **kwargs):
        super().__init__(
            master, bg=INPUT_BG, fg=TEXT_MAIN, insertbackground=TEXT_MAIN, relief='flat',
            highlightthickness=1, highlightbackground=INPUT_BORDER, highlightcolor=INPUT_BORDER_FOCUS,
            font=FONT_ENTRY, width=width, **kwargs
        )
        self.placeholder = placeholder
        self.placeholder_color = PLACEHOLDER_COLOR
        self.default_fg = TEXT_MAIN
        self._has_placeholder = False
        self._is_focused = False
        self.bind('<FocusIn>', self._on_focus_in)
        self.bind('<FocusOut>', self._on_focus_out)
        self.bind('<KeyRelease>', self._on_key_release)
        self._add_placeholder()

    def _on_focus_in(self, event=None):
        self._is_focused = True
        if self._has_placeholder:
            self.delete(0, 'end')
            self.config(fg=self.default_fg)
            self._has_placeholder = False

    def _on_focus_out(self, event=None):
        self._is_focused = False
        if not self.get().strip():
            self._add_placeholder()

    def _on_key_release(self, event=None):
        if not self._is_focused and not self.get().strip():
            self._add_placeholder()

    def _add_placeholder(self):
        self.delete(0, 'end')
        self.insert(0, self.placeholder)
        self.config(fg=self.placeholder_color)
        self._has_placeholder = True

    def get_value(self):
        if self._has_placeholder:
            return ""
        return self.get().strip()

    def reset(self):
        self._add_placeholder()

    def is_filled(self):
        return bool(self.get_value())


class NumericEntry(PlaceholderEntry):
    def __init__(self, master=None, placeholder="", width=None, **kwargs):
        super().__init__(master, placeholder, width, **kwargs)
        self.bind('<KeyRelease>', self._validate_input)

    def _validate_input(self, event):
        if event.keysym in ('BackSpace', 'Delete', 'Tab', 'Escape', 'Left', 'Right', 'Home', 'End', 'Return'):
            return None
        if event.state & 4:
            return None
        if event.char and not event.char.isdigit():
            return 'break'
        return None


class PlaceholderCombobox(ttk.Combobox):
    def __init__(self, master=None, values=None, default="New", placeholder="Select status", width=11, **kwargs):
        values = list(values) if values else ["New", "Revise"]
        self.actual_values = [v for v in values if v != placeholder]
        self.placeholder = placeholder
        self.default_value = default if default in self.actual_values else (self.actual_values[0] if self.actual_values else placeholder)
        self.placeholder_color = PLACEHOLDER_COLOR
        self.default_fg = TEXT_MAIN
        super().__init__(
            master, values=self.actual_values, state='readonly', width=width,
            style='Aurora.TCombobox', **kwargs
        )
        self.set(self.default_value)
        self.bind('<<ComboboxSelected>>', self._on_select)
        self.bind('<FocusOut>', self._on_focus_out)

    def _on_select(self, event=None):
        if self.get() in self.actual_values:
            self.config(foreground=self.default_fg)
        else:
            self.set(self.default_value)
            self.config(foreground=self.default_fg)

    def _on_focus_out(self, event=None):
        if self.get() not in self.actual_values:
            self.set(self.default_value)
            self.config(foreground=self.default_fg)

    def get_value(self):
        return self.get() if self.get() in self.actual_values else ""

    def reset(self):
        self.set(self.default_value)
        self.config(foreground=self.default_fg)

    def is_filled(self):
        return self.get() in self.actual_values


# ==========================================
# EFLAPP MAIN CLASS
# ==========================================
class EFLApp:
    def __init__(self, root, container=None, standalone=True, profile=None, on_open_settings=None):
        self.root = container if container is not None else root
        self.standalone = standalone and container is None
        self.on_open_settings = on_open_settings
        if profile is not None and profile != (None, None):
            self.user_id, self.user_code = profile
        else:
            saved = load_saved_profile()
            if saved:
                self.user_id, self.user_code = saved
            else:
                self.user_id = USER_ID
                self.user_code = USER_CODE

        self.selected_date = datetime.now()
        self.gs_manager = None
        self._anim_running = True
        self._updating = False
        self._msg_after_id = None
        self.headers1 = []
        self.headers2 = []
        self._refresh_lock = threading.Lock()
        self._submit_lock = threading.Lock()
        self._is_past_date = False

        self._sheet1_cache = []
        self._sheet2_cache = []
        self._cache_time = 0
        self._cache_ttl = 5

        self._auto_refresh_interval = 60
        self._auto_refresh_id = None

        self._setup_window()
        self._setup_styles()
        self._build_ui()

        if self.standalone:
            self.root.after(10, self._apply_dark_title_bar)
        self.root.after(50, self._init_google_sheets_async)

    def set_profile(self, profile):
        old_user = self.user_id
        if profile and profile != (None, None):
            self.user_id, self.user_code = profile
        else:
            self.user_id = USER_ID
            self.user_code = USER_CODE

        if hasattr(self, 'user_label') and self.user_label.winfo_exists():
            self.user_label.config(text=self.user_id or "Not configured")
        if hasattr(self, 'user_code_label') and self.user_code_label.winfo_exists():
            self.user_code_label.config(text=self.user_code or "")

        self._cache_time = 0
        if self.user_id != old_user:
            self._reset_all_rows()
        if self.gs_manager and self.gs_manager.connected:
            self._refresh_all_counts()

    def close(self):
        self._anim_running = False
        for cb_id in (self._auto_refresh_id, self._msg_after_id):
            if cb_id:
                try:
                    self.root.after_cancel(cb_id)
                except Exception:
                    pass
        self._auto_refresh_id = None
        self._msg_after_id = None

    def _apply_dark_title_bar(self):
        if sys.platform != "win32":
            return
        try:
            self.root.update_idletasks()
            hwnd = ctypes.windll.user32.GetParent(self.root.winfo_id())
            for attr in (20, 19):
                value = ctypes.c_int(1)
                result = ctypes.windll.dwmapi.DwmSetWindowAttribute(
                    hwnd, attr, ctypes.byref(value), ctypes.sizeof(value))
                if result == 0:
                    print(f"Dark title bar applied (attribute {attr})")
                    break
        except Exception as e:
            print(f"Dark title bar error: {e}")

    def _show_message(self, text, color=MSG_INFO, duration=3000):
        try:
            if hasattr(self, 'msg_label') and self.msg_label.winfo_exists():
                self.msg_label.config(text=text, fg=color)
                if hasattr(self, 'msg_icon') and self.msg_icon.winfo_exists():
                    icon = '✓' if color == MSG_SUCCESS else ('✕' if color == MSG_ERROR else '✦')
                    self.msg_icon.config(text=icon, fg=color)
                if self._msg_after_id:
                    self.root.after_cancel(self._msg_after_id)
                self._msg_after_id = self.root.after(duration, self._clear_message)
        except Exception:
            pass

    def _clear_message(self):
        try:
            if hasattr(self, 'msg_label') and self.msg_label.winfo_exists():
                self.msg_label.config(text='Ready', fg=TEXT_MUTED)
                if hasattr(self, 'msg_icon') and self.msg_icon.winfo_exists():
                    self.msg_icon.config(text='✦', fg=AURORA_CYAN)
            self._msg_after_id = None
        except Exception:
            pass

    def _init_google_sheets_async(self):
        self._show_message("Connecting to Google Sheets...", MSG_INFO, 60000)
        threading.Thread(target=self._connect_thread, daemon=True).start()

    def _connect_thread(self):
        try:
            self.gs_manager = GoogleSheetsManager(SPREADSHEET_URL)
            if self.gs_manager.connect():
                self.root.after(0, self._on_connected)
            else:
                self.root.after(0, self._on_disconnected)
        except Exception as e:
            print(f"Connection thread error: {e}")
            self.root.after(0, self._on_disconnected)

    def _on_connected(self):
        self._update_status(True)
        self._show_message("Connected to Google Sheets!", MSG_SUCCESS, 2500)
        self.root.after(50, self._refresh_all_counts)
        self._start_auto_refresh()

    def _on_disconnected(self):
        self._update_status(False)
        self._show_message("Failed to connect to Google Sheets", MSG_ERROR, 5000)

    def _update_status(self, connected):
        try:
            if hasattr(self, 'status_indicator') and hasattr(self, 'status_text'):
                if connected:
                    self.status_indicator.config(fg=STATUS_PILL_ONLINE_FG)
                    self.status_text.config(text="Online", fg=STATUS_PILL_ONLINE_FG)
                else:
                    self.status_indicator.config(fg=STATUS_PILL_OFFLINE_FG)
                    self.status_text.config(text="Offline", fg=STATUS_PILL_OFFLINE_FG)
        except Exception:
            pass
        if self.gs_manager:
            self.gs_manager.connected = connected
        self.root.update_idletasks()

    def _start_auto_refresh(self):
        self._schedule_next_auto_refresh()

    def _schedule_next_auto_refresh(self):
        if self._auto_refresh_id:
            try:
                self.root.after_cancel(self._auto_refresh_id)
            except Exception:
                pass
        self._auto_refresh_id = self.root.after(
            self._auto_refresh_interval * 1000, self._auto_refresh_tick)

    def _auto_refresh_tick(self):
        try:
            if self.gs_manager and self.gs_manager.connected:
                self._cache_time = 0
                self._refresh_all_counts()
                print("[AUTO-REFRESH] Counts updated")
            else:
                print("[AUTO-REFRESH] Reconnecting...")
                self._init_google_sheets_async()
        except Exception as e:
            print(f"[AUTO-REFRESH] Error: {e}")
        finally:
            self._schedule_next_auto_refresh()

    def _setup_window(self):
        if self.standalone:
            self.root.title("User KPI - EFL NEXUS")
            self.root.geometry("880x720")
            self.root.minsize(800, 650)
            self.root.configure(bg=BASE_BG)
        else:
            self.root.configure(bg=BASE_BG)
        self._widget_cache = {}

    def _setup_styles(self):
        style = ttk.Style()
        style.theme_use('clam')
        style.configure(
            'Aurora.TCombobox',
            fieldbackground=INPUT_BG,
            background=INPUT_BG,
            foreground=TEXT_MAIN,
            arrowcolor=TEXT_MUTED,
            relief='flat',
            padding=4
        )
        style.map(
            'Aurora.TCombobox',
            fieldbackground=[('readonly', INPUT_BG)],
            foreground=[('readonly', TEXT_MAIN)],
            selectbackground=[('readonly', INPUT_BG)],
            selectforeground=[('readonly', TEXT_MAIN)]
        )
        self.root.option_add('*TCombobox*Listbox.background', '#ffffff')
        self.root.option_add('*TCombobox*Listbox.foreground', TEXT_MAIN)
        self.root.option_add('*TCombobox*Listbox.selectBackground', '#0b1420')
        self.root.option_add('*TCombobox*Listbox.selectForeground', '#ffffff')

    def _build_ui(self):
        self._build_header()
        self._build_user_row()
        self._build_top_panel()
        self._build_bottom_panel()
        self._build_footer()
        self.root.update_idletasks()

    def _build_header(self):
        self.header = tk.Frame(self.root, bg=ACCENT_DARK)
        self.header.pack(fill='x', side='top')

        header_inner = tk.Frame(self.header, bg=ACCENT_DARK, padx=20, pady=12)
        header_inner.pack(fill='x')

        left_box = tk.Frame(header_inner, bg=ACCENT_DARK)
        left_box.pack(side='left', fill='y')

        title_row = tk.Frame(left_box, bg=ACCENT_DARK)
        title_row.pack(anchor='w')

        self.header_label = tk.Label(
            title_row, text='User KPI', bg=ACCENT_DARK, fg=TEXT_WHITE, font=FONT_HEADER
        )
        self.header_label.pack(side='left')

        tk.Label(
            left_box,
            text='Operator task logging, job record management, Google Sheets live sync, and daily KPI tracking',
            bg=ACCENT_DARK, fg=TEXT_LIGHT, font=('Segoe UI', 9)
        ).pack(anchor='w', pady=(2, 0))

        right_box = tk.Frame(header_inner, bg=ACCENT_DARK)
        right_box.pack(side='right', fill='y')

        if self.on_open_settings:
            settings_btn = tk.Label(
                right_box, text="⚙ Settings", bg=HEADER_BTN_BG, fg=TEXT_WHITE,
                font=('Segoe UI', 8, 'bold'), cursor='hand2', padx=8, pady=4
            )
            settings_btn.pack(side='right', padx=(8, 0))
            settings_btn.bind('<Button-1>', lambda e: self.on_open_settings())
            settings_btn.bind('<Enter>', lambda e, b=settings_btn: b.config(bg=HEADER_BTN_HOVER))
            settings_btn.bind('<Leave>', lambda e, b=settings_btn: b.config(bg=HEADER_BTN_BG))

    def _animate_header(self):
        pass

    def _build_user_row(self):
        row = tk.Frame(self.root, bg=BASE_BG)
        row.pack(fill='x', padx=18, pady=(10, 8))

        user_chip = tk.Frame(row, bg='#ffffff', highlightbackground=CARD_BORDER, highlightthickness=1, padx=10, pady=4)
        user_chip.pack(side='left')
        user_icon = tk.Label(user_chip, text='👤', bg='#ffffff', fg=TEXT_MUTED, font=('Segoe UI', 10))
        user_icon.pack(side='left', padx=(0, 5))
        tk.Label(user_chip, text='Operator:', bg='#ffffff', fg=TEXT_MUTED, font=('Segoe UI', 9)).pack(side='left', padx=(0, 4))
        self.user_label = tk.Label(user_chip, text=self.user_id or 'Not configured', bg='#ffffff', fg=TEXT_MAIN, font=FONT_LABEL)
        self.user_label.pack(side='left')

        # Keep User Code prominently as requested
        tk.Label(user_chip, text='|', bg='#ffffff', fg=CARD_BORDER, font=('Segoe UI', 9)).pack(side='left', padx=6)
        tk.Label(user_chip, text='Code:', bg='#ffffff', fg=TEXT_MUTED, font=('Segoe UI', 9)).pack(side='left', padx=(0, 4))
        self.user_code_label = tk.Label(user_chip, text=self.user_code or '', bg='#ffffff', fg=TEXT_MAIN, font=FONT_LABEL)
        self.user_code_label.pack(side='left')

        date_bar = tk.Frame(row, bg=BASE_BG)
        date_bar.pack(side='right')

        self.date_mode_label = tk.Label(date_bar, text="", bg=BASE_BG, font=('Segoe UI', 9, 'bold'))
        self.date_mode_label.pack(side='left', padx=(0, 8))

        tk.Label(date_bar, text='Date:', bg=BASE_BG, fg=TEXT_MUTED, font=('Segoe UI', 9, 'bold')).pack(side='left', padx=(0, 6))
        self.date_picker = DatePicker(date_bar, callback=self.on_date_selected)
        self.date_picker.frame.pack(side='left')

        self.today_btn = tk.Button(
            date_bar,
            text='Today',
            bg='#ffffff',
            fg=TEXT_MAIN,
            activebackground='#f1f5f9',
            activeforeground=TEXT_MAIN,
            relief='flat',
            highlightbackground=CARD_BORDER,
            highlightthickness=1,
            font=('Segoe UI', 9, 'bold'),
            cursor='hand2',
            padx=8,
            pady=2,
            command=lambda: self.date_picker.go_to_today()
        )
        self.today_btn.pack(side='left', padx=(8, 0))

    def on_date_selected(self, date_obj):
        self.selected_date = date_obj
        today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
        selected = date_obj.replace(hour=0, minute=0, second=0, microsecond=0)
        self._is_past_date = selected < today
        if self._is_past_date:
            self._show_message("Past date selected (View Only)", MSG_WARNING, 3000)
            self.date_mode_label.config(text="VIEW ONLY", fg=MSG_ERROR)
        else:
            self._clear_message()
            self.date_mode_label.config(text="", fg=TEXT_MAIN)
        self._cache_time = 0
        self.root.after(50, self._refresh_all_counts)

    def _make_row(self, parent, label_text, has_extra_counts=False, has_lp_count=False, dropdown_options=None):
        row = tk.Frame(parent, bg=CARD_BG)
        row.pack(fill='x', padx=6, pady=4)

        lbl = tk.Label(row, text=label_text, bg=CARD_BG, fg=TEXT_MAIN, font=FONT_LABEL, width=18, anchor='w')
        lbl.pack(side='left')

        right_frame = tk.Frame(row, bg=CARD_BG)
        right_frame.pack(side='left', padx=(10, 0), fill='x', expand=True)

        entry = PlaceholderEntry(right_frame, placeholder='Job ID', width=14)
        entry.pack(side='left', padx=(0, 6), ipady=3)

        extra_entries = []
        combo = None

        if has_extra_counts:
            e1 = NumericEntry(right_frame, placeholder='Load ID Count', width=12)
            e1.pack(side='left', padx=(0, 6), ipady=3)
            extra_entries.append(e1)

            if has_lp_count:
                e2 = NumericEntry(right_frame, placeholder='LP Count', width=12)
                e2.pack(side='left', padx=(0, 6), ipady=3)
                extra_entries.append(e2)
        elif dropdown_options:
            combo = PlaceholderCombobox(right_frame, values=dropdown_options, default="New", width=11)
            combo.pack(side='left', padx=(0, 6), ipady=2)

        submit_btn = tk.Button(
            right_frame,
            text='Submit',
            bg=BTN_SUBMIT_BG,
            fg=TEXT_WHITE,
            activebackground=BTN_SUBMIT_HOVER,
            activeforeground=TEXT_WHITE,
            relief='flat',
            font=FONT_BUTTON,
            cursor='hand2',
            width=7,
            command=lambda: self._on_submit(label_text, entry, extra_entries if has_extra_counts else None, combo)
        )
        submit_btn.pack(side='left', padx=3)

        delete_btn = tk.Button(
            right_frame,
            text='Delete',
            bg=BTN_DELETE_BG,
            fg=TEXT_WHITE,
            activebackground=BTN_DELETE_HOVER,
            activeforeground=TEXT_WHITE,
            relief='flat',
            font=FONT_BUTTON,
            cursor='hand2',
            width=7,
            command=lambda: self._on_delete(label_text, entry, extra_entries if has_extra_counts else None, combo)
        )
        delete_btn.pack(side='left', padx=3)

        count_frame = tk.Frame(right_frame, bg='#f8fafc', highlightbackground=CARD_BORDER, highlightthickness=1, padx=4, pady=2)
        count_frame.pack(side='left', padx=(8, 0))

        count_label = tk.Label(count_frame, text='0', bg='#f8fafc', fg=TEXT_MAIN, font=FONT_COUNT, width=7)
        count_label.pack()

        return {
            'entry': entry,
            'extra_entries': extra_entries,
            'count_label': count_label,
            'combo': combo,
            'label': label_text,
            'submit_btn': submit_btn,
            'delete_btn': delete_btn,
        }

    def _build_top_panel(self):
        card = tk.Frame(self.root, bg=CARD_BG, highlightbackground=CARD_BORDER, highlightthickness=1, padx=16, pady=12)
        card.pack(fill='x', padx=18, pady=(0, 10))

        header = tk.Frame(card, bg=CARD_BG)
        header.pack(fill='x', pady=(0, 10))

        tk.Label(header, text="Reconciliation & Load Plan", bg=CARD_BG, fg=TEXT_MAIN, font=('Segoe UI', 11, 'bold')).pack(side='left')
        tk.Label(header, text="— GDN, GRN, and Load Plan entry", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 9)).pack(side='left', padx=(8, 0))

        self.rows_top = {}
        top_configs = [
            ('GDN Reconciliation:', ['New', 'Revise']),
            ('GRN Reconciliation:', ['New', 'Revise']),
            ('GDN Creation:', ['New', 'Revise', 'Separate']),
            ('GRN Creation:', ['New', 'Revise', 'Separate']),
            ('Load Plan or Asn:', ['New', 'Revise'])
        ]

        for label, options in top_configs:
            self.rows_top[label] = self._make_row(card, label, has_extra_counts=False, dropdown_options=options)

        self.root.update_idletasks()

    def _build_bottom_panel(self):
        card = tk.Frame(self.root, bg=CARD_BG, highlightbackground=CARD_BORDER, highlightthickness=1, padx=16, pady=12)
        card.pack(fill='x', padx=18, pady=(0, 10))

        header = tk.Frame(card, bg=CARD_BG)
        header.pack(fill='x', pady=(0, 10))

        tk.Label(header, text="Audit, Shipping & Transfers", bg=CARD_BG, fg=TEXT_MAIN, font=('Segoe UI', 11, 'bold')).pack(side='left')
        tk.Label(header, text="— Load audit, shipping confirmation, transfer movements & allocation", bg=CARD_BG, fg=TEXT_MUTED, font=('Segoe UI', 9)).pack(side='left', padx=(8, 0))

        self.rows_bottom = {}
        row_configs = [
            ('Load Audit:', False),
            ('Shipping:', False),
            ('Load Transfer:', True),
            ('Allocation/Backorders:', True)
        ]

        for label, has_lp in row_configs:
            self.rows_bottom[label] = self._make_row(card, label, has_extra_counts=True, has_lp_count=has_lp)

        self.root.update_idletasks()

    def _build_footer(self):
        footer = tk.Frame(self.root, bg=BASE_BG)
        footer.pack(fill='x', padx=18, pady=(0, 12))

        top_row = tk.Frame(footer, bg=BASE_BG)
        top_row.pack(fill='x')

        msg_frame = tk.Frame(top_row, bg='#ffffff', highlightbackground=CARD_BORDER, highlightthickness=1, height=32, padx=10)
        msg_frame.pack(side='left', fill='x', expand=True, padx=(0, 12))
        msg_frame.pack_propagate(False)

        self.msg_icon = tk.Label(msg_frame, text='✦', bg='#ffffff', fg=AURORA_CYAN, font=('Segoe UI', 9, 'bold'))
        self.msg_icon.pack(side='left', padx=(0, 6))

        self.msg_label = tk.Label(msg_frame, text='Ready', bg='#ffffff', fg=TEXT_MUTED, font=FONT_MSG, anchor='w')
        self.msg_label.pack(side='left', fill='x', expand=True)

        btn_frame = tk.Frame(top_row, bg=BASE_BG)
        btn_frame.pack(side='right')

        self.download_btn = tk.Button(
            btn_frame,
            text='Download KPI',
            bg='#0284c7',
            fg='#ffffff',
            activebackground='#0369a1',
            activeforeground='#ffffff',
            relief='flat',
            font=('Segoe UI', 9, 'bold'),
            cursor='hand2',
            padx=12,
            pady=4,
            command=self._on_download
        )
        self.download_btn.pack(side='right', padx=(6, 0))

        self.refresh_btn = tk.Button(
            btn_frame,
            text='Refresh',
            bg='#0d1b2a',
            fg='#ffffff',
            activebackground='#162e4c',
            activeforeground=AURORA_CYAN,
            relief='flat',
            font=('Segoe UI', 9, 'bold'),
            cursor='hand2',
            padx=12,
            pady=4,
            command=self._on_refresh
        )
        self.refresh_btn.pack(side='right')

        self.root.update_idletasks()

    @staticmethod
    def _parse_timestamp_date(ts):
        if not ts:
            return ""
        ts = str(ts).strip()
        formats = ["%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y",
                   "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y",
                   "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"]
        for fmt in formats:
            try:
                return datetime.strptime(ts, fmt).strftime("%d-%m-%Y")
            except ValueError:
                continue
        return ts.split(" ")[0] if " " in ts else ts

    def _get_task_count_from_values_sheet1(self, all_values, task_name, user_id, date_str):
        if not all_values or len(all_values) <= 1:
            return 0
        task_norm_target = task_name.replace(":", "").strip().lower()
        user_target = user_id.strip().lower()
        seen = set()
        count = 0
        for row in all_values[1:]:
            if not row or len(row) < 5:
                continue
            row_timestamp = row[0].strip() if row[0] else ""
            row_task = row[1].strip() if row[1] else ""
            row_user = row[4].strip() if row[4] else ""
            row_date = self._parse_timestamp_date(row_timestamp)
            row_task_norm = row_task.replace(":", "").strip().lower()
            if row_task_norm != task_norm_target:
                continue
            if row_user.lower() != user_target:
                continue
            if date_str and row_date != date_str:
                continue
            key = f"{row_timestamp}|{row_user.lower()}|{row_task_norm}"
            if key not in seen:
                seen.add(key)
                count += 1
        return count

    def _get_sum_from_values_sheet2(self, all_values, task_name, user_id, date_str):
        if not all_values or len(all_values) <= 1:
            return 0, 0
        task_norm_target = task_name.replace(":", "").strip().lower()
        user_target = user_id.strip().lower()
        seen = set()
        load_sum = 0
        lp_sum = 0
        for row in all_values[1:]:
            if not row or len(row) < 6:
                continue
            row_timestamp = row[0].strip() if row[0] else ""
            row_task = row[1].strip() if row[1] else ""
            row_user = row[5].strip() if row[5] else ""
            load_val = row[3].strip() if len(row) > 3 and row[3] else ""
            lp_val = row[4].strip() if len(row) > 4 and row[4] else ""
            row_date = self._parse_timestamp_date(row_timestamp)
            row_task_norm = row_task.replace(":", "").strip().lower()
            if row_task_norm != task_norm_target:
                continue
            if row_user.lower() != user_target:
                continue
            if date_str and row_date != date_str:
                continue
            key = f"{row_timestamp}|{row_user.lower()}|{row_task_norm}"
            if key in seen:
                continue
            seen.add(key)
            try:
                load_sum += int(float(load_val)) if load_val else 0
            except ValueError:
                pass
            try:
                lp_sum += int(float(lp_val)) if lp_val else 0
            except ValueError:
                pass
        return load_sum, lp_sum

    def _is_duplicate_in_sheet1(self, task_name, job_id, job_status):
        if not self.gs_manager or not self.gs_manager.connected:
            return None
        try:
            all_values = self.gs_manager.get_all_values_sheet1()
            self._sheet1_cache = all_values
            if not all_values or len(all_values) <= 1:
                return None

            def _norm(t):
                return t.replace(":", "").replace(" ", "").strip().lower()

            task_groups = [
                {"gdncreation", "gdnreconciliation"},
                {"grncreation", "grnreconciliation"},
            ]
            submitted_norm = _norm(task_name)
            matched_group = None
            for grp in task_groups:
                if submitted_norm in grp:
                    matched_group = grp
                    break
            job_norm = str(job_id).strip().lower()
            status_norm = str(job_status).strip().lower()

            for idx, row in enumerate(all_values[1:], start=2):
                if not row or len(row) < 5:
                    continue
                row_task = row[1].strip() if len(row) > 1 and row[1] else ""
                row_job = row[2].strip() if len(row) > 2 and row[2] else ""
                row_status = row[3].strip() if len(row) > 3 and row[3] else ""
                row_user = row[4].strip() if len(row) > 4 and row[4] else ""
                row_task_norm = _norm(row_task)
                if row_job.lower() != job_norm:
                    continue
                if row_status.lower() != status_norm:
                    continue
                matched = False
                if matched_group is not None:
                    if row_task_norm in matched_group:
                        matched = True
                else:
                    if row_task_norm == submitted_norm:
                        matched = True
                if matched:
                    return row_user if row_user else "another user"
            return None
        except Exception as e:
            print(f"[DUP-CHECK] Error: {e}")
            return None

    def _get_owner_of_row_sheet1(self, job_id):
        if not self.gs_manager or not self.gs_manager.connected:
            return None
        try:
            all_values = self.gs_manager.get_all_values_sheet1()
            self._sheet1_cache = all_values
            if not all_values or len(all_values) <= 1:
                return None
            job_norm = str(job_id).strip().lower()
            for row in all_values[1:]:
                if not row or len(row) < 3:
                    continue
                row_job = row[2].strip() if len(row) > 2 and row[2] else ""
                if row_job.lower() == job_norm:
                    row_user = row[4].strip() if len(row) > 4 and row[4] else ""
                    return row_user if row_user else "another user"
            return None
        except Exception as e:
            print(f"Owner lookup error (Sheet1): {e}")
            return None

    def _get_owner_of_row_sheet2(self, job_id):
        if not self.gs_manager or not self.gs_manager.connected:
            return None
        try:
            all_values = self.gs_manager.get_all_values_sheet2()
            self._sheet2_cache = all_values
            if not all_values or len(all_values) <= 1:
                return None
            job_norm = str(job_id).strip().lower()
            for row in all_values[1:]:
                if not row or len(row) < 3:
                    continue
                row_job = row[2].strip() if len(row) > 2 and row[2] else ""
                if row_job.lower() == job_norm:
                    row_user = row[5].strip() if len(row) > 5 and row[5] else ""
                    return row_user if row_user else "another user"
            return None
        except Exception as e:
            print(f"Owner lookup error (Sheet2): {e}")
            return None

    def _refresh_all_counts(self):
        if self._updating or not self.gs_manager or not self.gs_manager.connected:
            return
        threading.Thread(target=self._refresh_counts_thread, daemon=True).start()

    def _refresh_counts_thread(self):
        with self._refresh_lock:
            if self._updating:
                return
            self._updating = True
            try:
                now = time.time()
                use_cache = (now - self._cache_time) < self._cache_ttl
                if use_cache and self._sheet1_cache and self._sheet2_cache:
                    values1 = self._sheet1_cache
                    values2 = self._sheet2_cache
                else:
                    result = {}
                    def fetch_sheet1():
                        try:
                            result['v1'] = self.gs_manager.get_all_values_sheet1()
                        except Exception as e:
                            result['v1'] = []
                    def fetch_sheet2():
                        try:
                            result['v2'] = self.gs_manager.get_all_values_sheet2()
                        except Exception as e:
                            result['v2'] = []
                    t1 = threading.Thread(target=fetch_sheet1, daemon=True)
                    t2 = threading.Thread(target=fetch_sheet2, daemon=True)
                    t1.start(); t2.start()
                    t1.join(timeout=15); t2.join(timeout=15)
                    values1 = result.get('v1', [])
                    values2 = result.get('v2', [])
                    self._sheet1_cache = values1
                    self._sheet2_cache = values2
                    self._cache_time = time.time()

                selected_date = self.selected_date
                user_id = self.user_id
                self.root.after(0, lambda: self._update_counts_ui(values1, values2, selected_date, user_id))
            except Exception as e:
                print(f"Error refreshing counts: {e}")
            finally:
                self._updating = False

    def _update_counts_ui(self, values1, values2, selected_date, user_id):
        try:
            date_str = selected_date.strftime('%d-%m-%Y') if hasattr(selected_date, 'strftime') else str(selected_date)

            combined_pairs = {
                "GDN Creation:": ["GDN Creation:", "GDN Reconciliation:"],
                "GRN Creation:": ["GRN Creation:", "GRN Reconciliation:"],
            }

            for label, row_data in self.rows_top.items():
                if label in combined_pairs:
                    total = 0
                    for part in combined_pairs[label]:
                        total += self._get_task_count_from_values_sheet1(values1, part, user_id, date_str)
                    row_data['count_label'].config(text=str(total))
                else:
                    cnt = self._get_task_count_from_values_sheet1(values1, label, user_id, date_str)
                    row_data['count_label'].config(text=str(cnt))

            for label, row_data in self.rows_bottom.items():
                l_sum, lp_sum = self._get_sum_from_values_sheet2(values2, label, user_id, date_str)
                if label in ('Load Transfer:', 'Allocation/Backorders:'):
                    txt = f"{l_sum}/{lp_sum}"
                else:
                    txt = str(l_sum)
                row_data['count_label'].config(text=txt)

            self.root.update_idletasks()
        except Exception as e:
            print(f"Error updating UI counts: {e}")

    def _reset_all_rows(self):
        if self._is_past_date:
            return
        for row_data in self.rows_top.values():
            if 'entry' in row_data and row_data['entry']:
                row_data['entry'].reset()
            if 'combo' in row_data and row_data['combo']:
                row_data['combo'].reset()
        for row_data in self.rows_bottom.values():
            if 'entry' in row_data and row_data['entry']:
                row_data['entry'].reset()
            if 'extra_entries' in row_data and row_data['extra_entries']:
                for entry in row_data['extra_entries']:
                    if entry:
                        entry.reset()
        self._show_message("All rows reset", MSG_INFO, 1500)
        self.root.after(100, self._refresh_all_counts)

    def _validate_row_fields(self, row_data, is_bottom_panel=False):
        if self._is_past_date:
            return False, ["Date is in the past - View Only mode"]
        missing_fields = []
        if 'entry' in row_data and row_data['entry']:
            if not row_data['entry'].is_filled():
                missing_fields.append("Job ID")
        else:
            missing_fields.append("Job ID")
        if not is_bottom_panel:
            if 'combo' in row_data and row_data['combo']:
                if not row_data['combo'].is_filled():
                    missing_fields.append("Job Status")
            else:
                missing_fields.append("Job Status")
        if is_bottom_panel and 'extra_entries' in row_data and row_data['extra_entries']:
            if len(row_data['extra_entries']) > 0 and row_data['extra_entries'][0]:
                if not row_data['extra_entries'][0].is_filled():
                    missing_fields.append("Load ID Count")
            else:
                missing_fields.append("Load ID Count")
            if len(row_data['extra_entries']) > 1 and row_data['extra_entries'][1]:
                if not row_data['extra_entries'][1].is_filled():
                    missing_fields.append("LP Count")
        return len(missing_fields) == 0, missing_fields

    def _save_to_sheet1(self, section, data):
        if self._is_past_date:
            self._show_message("Cannot save - Past date (View Only)", MSG_WARNING, 2000)
            return False
        if not self.gs_manager or not self.gs_manager.connected:
            self._show_message("Offline - Cannot save", MSG_ERROR, 2000)
            return False
        try:
            timestamp = datetime.now().strftime("%d-%m-%Y %H:%M:%S")
            job_status = data.get("type", "New")
            task_name = section
            row = [timestamp, task_name, data.get("job_id", ""), job_status, self.user_id, self.user_code]
            success = self.gs_manager.append_row_sheet1(row)
            if success:
                if enqueue_job and is_reconciliation_task(task_name):
                    try:
                        enqueue_job(data.get("job_id", ""), task_name, self.user_id)
                    except Exception:
                        pass
                self._show_message(f"Saved: {data.get('job_id', '')}", MSG_SUCCESS, 2000)
                self._cache_time = 0
                self.root.after(100, self._refresh_all_counts)
                return True
            return False
        except Exception as e:
            print(f"Error saving to Sheet1: {e}")
            self._show_message("Save failed", MSG_ERROR, 2000)
            return False

    def _save_to_sheet2(self, section, data):
        if self._is_past_date:
            self._show_message("Cannot save - Past date (View Only)", MSG_WARNING, 2000)
            return False
        if not self.gs_manager or not self.gs_manager.connected:
            self._show_message("Offline - Cannot save", MSG_ERROR, 2000)
            return False
        try:
            timestamp = datetime.now().strftime("%d-%m-%Y %H:%M:%S")
            task_name = section
            row = [timestamp, task_name, data.get("job_id", ""),
                   data.get("load_count", ""), data.get("lp_count", ""), self.user_id, self.user_code]
            success = self.gs_manager.append_row_sheet2(row)
            if success:
                self._show_message(f"Saved: {data.get('job_id', '')}", MSG_SUCCESS, 2000)
                self._cache_time = 0
                self.root.after(100, self._refresh_all_counts)
                return True
            return False
        except Exception as e:
            print(f"Error saving to Sheet2: {e}")
            self._show_message("Save failed", MSG_ERROR, 2000)
            return False

    def _delete_from_sheet1(self, job_id):
        if self._is_past_date:
            self._show_message("Cannot delete - Past date (View Only)", MSG_WARNING, 2000)
            return False
        if not self.gs_manager or not self.gs_manager.connected:
            self._show_message("Offline - Cannot delete", MSG_ERROR, 2000)
            return False
        try:
            owner = self._get_owner_of_row_sheet1(job_id)
            if owner is None:
                self._show_message("Job ID not found", MSG_ERROR, 2000)
                return False
            if owner.lower() != self.user_id.lower():
                self._show_message(f"Cannot delete - Assigned by {owner}", MSG_WARNING, 3000)
                return False
            success = self.gs_manager.search_and_delete_row_sheet1(2, job_id, self.user_id)
            if success:
                self._show_message(f"Deleted: {job_id}", MSG_SUCCESS, 2000)
                self._cache_time = 0
                self.root.after(100, self._refresh_all_counts)
                return True
            return False
        except Exception as e:
            print(f"Error deleting from Sheet1: {e}")
            self._show_message("Delete failed", MSG_ERROR, 2000)
            return False

    def _delete_from_sheet2(self, job_id):
        if self._is_past_date:
            self._show_message("Cannot delete - Past date (View Only)", MSG_WARNING, 2000)
            return False
        if not self.gs_manager or not self.gs_manager.connected:
            self._show_message("Offline - Cannot delete", MSG_ERROR, 2000)
            return False
        try:
            owner = self._get_owner_of_row_sheet2(job_id)
            if owner is None:
                self._show_message("Job ID not found", MSG_ERROR, 2000)
                return False
            if owner.lower() != self.user_id.lower():
                self._show_message(f"Cannot delete - Assigned by {owner}", MSG_WARNING, 3000)
                return False
            success = self.gs_manager.search_and_delete_row_sheet2(2, job_id, self.user_id)
            if success:
                self._show_message(f"Deleted: {job_id}", MSG_SUCCESS, 2000)
                self._cache_time = 0
                self.root.after(100, self._refresh_all_counts)
                return True
            return False
        except Exception as e:
            print(f"Error deleting from Sheet2: {e}")
            self._show_message("Delete failed", MSG_ERROR, 2000)
            return False

    def _on_submit(self, section, entry, extra_entries=None, combo=None):
        if self._is_past_date:
            self._show_message("Cannot submit - Past date (View Only)", MSG_WARNING, 2000)
            return
        is_bottom = extra_entries is not None
        row_data = {"entry": entry, "extra_entries": extra_entries, "combo": combo}
        is_valid, _ = self._validate_row_fields(row_data, is_bottom)
        if not is_valid:
            self._show_message("Fill all fields", MSG_ERROR, 2000)
            return
        job_id = entry.get_value()
        if not job_id:
            self._show_message("Enter Job ID", MSG_ERROR, 1500)
            return
        data = {"job_id": job_id, "type": combo.get_value() if combo else ""}
        if extra_entries:
            data["load_count"] = extra_entries[0].get_value() if len(extra_entries) > 0 else ""
            data["lp_count"] = extra_entries[1].get_value() if len(extra_entries) > 1 else ""
            if self._save_to_sheet2(section, data):
                self._reset_all_rows()
        else:
            job_status = combo.get_value() if combo else ""
            with self._submit_lock:
                existing_user = self._is_duplicate_in_sheet1(section, job_id, job_status)
                if existing_user:
                    self._show_message(f"Already assigned by {existing_user}", MSG_WARNING, 3000)
                    return
                if self._save_to_sheet1(section, data):
                    self._reset_all_rows()

    def _on_delete(self, section, entry, extra_entries=None, combo=None):
        if self._is_past_date:
            self._show_message("Cannot delete - Past date (View Only)", MSG_WARNING, 2000)
            return
        is_bottom = extra_entries is not None
        row_data = {"entry": entry, "extra_entries": extra_entries, "combo": combo}
        is_valid, _ = self._validate_row_fields(row_data, is_bottom)
        if not is_valid:
            self._show_message("Fill all fields to delete", MSG_ERROR, 2000)
            return
        job_id = entry.get_value()
        if not job_id:
            self._show_message("Enter Job ID", MSG_ERROR, 1500)
            return
        if extra_entries:
            success = self._delete_from_sheet2(job_id)
        else:
            success = self._delete_from_sheet1(job_id)
        if success:
            self._reset_all_rows()

    def _on_refresh(self):
        self._show_message("Refreshing...", MSG_INFO)
        self._cache_time = 0
        if not self._is_past_date:
            self._reset_all_rows()
        else:
            self._show_message("View Only - Refreshing counts", MSG_INFO, 1500)
        if self.gs_manager and self.gs_manager.connected:
            self._refresh_all_counts()
            self._show_message("Refreshed", MSG_SUCCESS, 1500)
        else:
            self._show_message("Reconnecting...", MSG_INFO, 3000)
            threading.Thread(target=self._reconnect_thread, daemon=True).start()
        self._schedule_next_auto_refresh()

    def _on_download(self):
        if not self.gs_manager or not self.gs_manager.connected:
            self._show_message("Offline - Cannot download", MSG_ERROR, 3000)
            return

        self._show_message("Downloading...", MSG_INFO, 5000)
        threading.Thread(target=self._do_download_thread, daemon=True).start()

    def _do_download_thread(self):
        try:
            values1 = self.gs_manager.get_all_values_sheet1()
            values2 = self.gs_manager.get_all_values_sheet2()
            self._sheet1_cache = values1
            self._sheet2_cache = values2
            self._cache_time = time.time()

            start_date, end_date = compute_export_range(self.selected_date)
            range_str = f"{start_date.strftime('%d-%m-%Y')} → {end_date.strftime('%d-%m-%Y')}"

            counts_by_day = build_counts_for_range(
                values1, values2, self.user_id, start_date, end_date,
                self._get_task_count_from_values_sheet1,
                self._get_sum_from_values_sheet2,
            )

            chosen_folder = [None]
            done = threading.Event()

            def pick_folder():
                chosen_folder[0] = filedialog.askdirectory(
                    title="Select folder to save the Excel file",
                    initialdir=os.path.dirname(os.path.abspath(__file__))
                )
                done.set()

            self.root.after(0, pick_folder)
            done.wait(timeout=180)

            if not chosen_folder[0]:
                self.root.after(0, lambda: self._show_message(
                    "Download cancelled — no folder selected", MSG_WARNING, 3000))
                return

            path = os.path.join(chosen_folder[0], EXCEL_FILENAME)

            if not ensure_excel_exists(path):
                self.root.after(0, lambda: self._show_message(
                    f"Could not create Excel at {path}", MSG_ERROR, 5000))
                return

            ok, message, details = write_counts_to_excel(
                path, self.user_id, self.user_code, counts_by_day
            )

            print("\n" + "=" * 60)
            print("[DOWNLOAD] Selected date:", self.selected_date.strftime('%d-%m-%Y'))
            print("[DOWNLOAD] Export range :", range_str)
            print("[DOWNLOAD] Folder       :", chosen_folder[0])
            print("[DOWNLOAD] File         :", path)
            for d in details:
                print("  " + d)
            print("=" * 60 + "\n")

            if ok:
                self.root.after(0, lambda: self._show_message(
                    f"✔ Excel saved ({range_str})", MSG_SUCCESS, 5000))
            else:
                self.root.after(0, lambda: self._show_message(
                    f"✘ {message}", MSG_ERROR, 5000))
        except Exception as e:
            print(f"[DOWNLOAD] Error: {e}")
            self.root.after(0, lambda: self._show_message(
                f"Download failed: {e}", MSG_ERROR, 5000))

    def _reconnect_thread(self):
        if not self.gs_manager:
            self.gs_manager = GoogleSheetsManager(SPREADSHEET_URL)
        ok = self.gs_manager.connect()
        self.root.after(0, lambda: self._on_reconnect_done(ok))

    def _on_reconnect_done(self, ok):
        self._update_status(ok)
        if ok:
            self._cache_time = 0
            self._refresh_all_counts()
            self._show_message("Connected & Refreshed", MSG_SUCCESS, 2000)
        else:
            self._show_message("Connection failed", MSG_ERROR, 3000)


if __name__ == "__main__":
    root = tk.Tk()
    app = EFLApp(root)
    root.mainloop()

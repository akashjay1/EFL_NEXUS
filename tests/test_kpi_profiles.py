import json
import tempfile
import tkinter as tk
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import Mock, patch

import efl_app
import main_app
from kpi_profile import export_filename, load_saved_profile, validate_profile
from outlook_email_gui import ConfigStore


class KPIProfileTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.users = self.folder / "efl_users.json"
        self.users.write_text(json.dumps({"users": {"Akash": {}, "Tharindu": {}}}), encoding="utf-8")
        self.config_path = self.folder / "config.json"

    def test_registered_name_is_canonical_and_code_is_required(self):
        self.assertEqual(validate_profile("  tharindu ", " C-22 ", self.users), ("Tharindu", "C-22"))
        with self.assertRaises(ValueError):
            validate_profile("Unknown", "C-22", self.users)
        with self.assertRaises(ValueError):
            validate_profile("Akash", " ", self.users)

    def test_profile_persists_and_missing_profile_stays_unconfigured(self):
        store = ConfigStore(str(self.config_path))
        self.assertIsNone(load_saved_profile(self.config_path, self.users))
        self.assertTrue(store.save(kpi_user_id="Tharindu", kpi_user_code="C-22"))
        self.assertEqual(load_saved_profile(self.config_path, self.users), ("Tharindu", "C-22"))
        self.assertEqual(ConfigStore(str(self.config_path)).config["kpi_user_id"], "Tharindu")

    def test_settings_save_rejects_invalid_and_updates_open_tool(self):
        app = main_app.MainApp.__new__(main_app.MainApp)
        app.config_store = ConfigStore(str(self.config_path))
        app.kpi_user_id_entry = Mock()
        app.kpi_user_code_entry = Mock()
        app.kpi_status_pill = Mock()
        app.kpi_msg_lbl = Mock()
        app.tool4_app = Mock()
        app.kpi_user_id_entry.get.return_value = "Nobody"
        app.kpi_user_code_entry.get.return_value = "C-22"
        with patch("main_app.validate_profile", side_effect=lambda u, c: validate_profile(u, c, self.users)):
            self.assertFalse(app._save_kpi_profile())
            app.kpi_user_id_entry.get.return_value = "tharindu"
            self.assertTrue(app._save_kpi_profile())
        app.tool4_app.set_profile.assert_called_once_with(("Tharindu", "C-22"))
        self.assertEqual(load_saved_profile(self.config_path, self.users), ("Tharindu", "C-22"))

    def test_settings_save_failure_does_not_activate_new_profile(self):
        app = main_app.MainApp.__new__(main_app.MainApp)
        app.config_store = Mock(config={"kpi_user_id": "Akash", "kpi_user_code": "OLD"})
        app.config_store.save.return_value = False
        app.kpi_user_id_entry = Mock()
        app.kpi_user_code_entry = Mock()
        app.kpi_user_id_entry.get.return_value = "Tharindu"
        app.kpi_user_code_entry.get.return_value = "NEW"
        app.kpi_msg_lbl = Mock()
        app.tool4_app = Mock()
        with patch("main_app.validate_profile", return_value=("Tharindu", "NEW")):
            self.assertFalse(app._save_kpi_profile())
        app.tool4_app.set_profile.assert_not_called()
        self.assertEqual(app.config_store.config["kpi_user_id"], "Akash")

    def test_tool_is_gated_until_profile_is_saved_and_switches_live(self):
        try:
            root = tk.Tk()
        except tk.TclError as error:
            self.skipTest(f"Tk unavailable: {error}")
        self.addCleanup(root.destroy)
        root.withdraw()
        page = tk.Frame(root)
        page.pack()
        with patch.object(efl_app.EFLApp, "_init_google_sheets_async"), patch("efl_app.load_saved_profile", return_value=None):
            tool = efl_app.EFLApp(root, container=page, standalone=False, profile=None)
            self.assertEqual(tool.download_btn.cget("state"), "disabled")
            self.assertTrue(tool._setup_overlay.winfo_exists())
            tool.set_profile(("Akash", "CSSUN151"))
            self.assertEqual(tool.download_btn.cget("state"), "normal")
            for row in tool.rows_top.values():
                combo = row["combo"]
                self.assertEqual(combo.cget("values"), ("New", "Revise"))
                self.assertEqual(combo.get(), "Select status")
                self.assertFalse(combo.is_filled())
            tool.gs_manager = Mock(connected=True)
            tool._refresh_all_counts = Mock()
            tool.set_profile(("Tharindu", "C-22"))
            self.assertEqual(tool.user_label.cget("text"), "Tharindu")
            tool._refresh_all_counts.assert_called_once()
            tool.close()


class KPIUserDataTests(unittest.TestCase):
    def setUp(self):
        self.app = efl_app.EFLApp.__new__(efl_app.EFLApp)
        self.app.user_id = "Tharindu"
        self.app.user_code = "C-22"
        self.app.selected_date = datetime(2026, 9, 3)
        self.app.gs_manager = Mock(connected=True)
        self.app._show_message = Mock()
        self.app._refresh_all_counts = Mock()
        self.app._get_owner_of_row_sheet1 = Mock(return_value="Akash")
        self.app._get_owner_of_row_sheet2 = Mock(return_value="Akash")

    @patch("efl_app.enqueue_job")
    def test_save_and_queue_use_active_user(self, enqueue_job):
        self.app.gs_manager.append_row_sheet1.return_value = True
        self.assertTrue(self.app._save_to_sheet1("GDN Reconciliation:", {"job_id": "JOB-1", "job_status": "New"}))
        self.assertEqual(self.app.gs_manager.append_row_sheet1.call_args.args[0][-1], "Tharindu")
        enqueue_job.assert_called_once_with("JOB-1", "GDN Reconciliation:", "Tharindu")
        self.app.gs_manager.append_row_sheet2.return_value = True
        self.assertTrue(self.app._save_to_sheet2("Shipping:", {"job_id": "JOB-2"}))
        self.assertEqual(self.app.gs_manager.append_row_sheet2.call_args.args[0][-1], "Tharindu")

    def test_delete_blocks_other_users_and_filters_by_active_user(self):
        self.app.gs_manager.search_and_delete_row_sheet1.return_value = False
        self.app.gs_manager.search_and_delete_row_sheet2.return_value = False
        self.assertFalse(self.app._delete_from_sheet1("JOB-1"))
        self.assertFalse(self.app._delete_from_sheet2("JOB-1"))
        self.app.gs_manager.search_and_delete_row_sheet1.assert_called_with(2, "JOB-1", "Tharindu")
        self.app.gs_manager.search_and_delete_row_sheet2.assert_called_with(2, "JOB-1", "Tharindu")
        self.app.gs_manager.search_and_delete_row_sheet1.return_value = True
        self.assertTrue(self.app._delete_from_sheet1("JOB-2"))
        self.app.gs_manager.search_and_delete_row_sheet1.assert_called_with(2, "JOB-2", "Tharindu")

    def test_counts_are_filtered_by_user(self):
        values = [["Timestamp", "Task", "Job ID", "Status", "User ID"],
                  ["01-09-2026 10:00:00", "GDN Creation:", "A", "New", "Akash"],
                  ["01-09-2026 11:00:00", "GDN Creation:", "B", "New", "Tharindu"]]
        self.assertEqual(self.app._get_task_count_from_values_sheet1(values, "GDN Creation:", "Tharindu", "01-09-2026"), 1)

    def test_download_requires_connection(self):
        self.app.gs_manager.connected = False
        with patch("efl_app.threading.Thread") as thread:
            self.app._on_download()
        thread.assert_not_called()

    def test_exports_are_separate_and_repeat_download_regenerates(self):
        with tempfile.TemporaryDirectory() as folder:
            header1 = ["Timestamp", "Task", "Job ID", "Status", "User ID"]
            header2 = ["Timestamp", "Task", "Job ID", "Load", "LP", "User ID"]
            values = [header1,
                      ["01-09-2026 10:00:00", "GDN Creation:", "A", "New", "Akash"],
                      ["01-09-2026 11:00:00", "GDN Creation:", "B", "New", "Tharindu"]]
            def export(user, code, data):
                return efl_app.generate_user_export(
                    folder, user, code, datetime(2026, 9, 3), data, [header2],
                    self.app._get_task_count_from_values_sheet1,
                    self.app._get_sum_from_values_sheet2,
                )
            ok_a, _, path_a, _ = export("Akash", "C-11", values)
            ok_t, _, path_t, _ = export("Tharindu", "C-22", values)
            self.assertTrue(ok_a and ok_t)
            self.assertNotEqual(path_a, path_t)
            self.assertEqual(Path(path_t).name, export_filename("Tharindu", datetime(2026, 9, 1), datetime(2026, 9, 2)))
            import openpyxl
            for path, user, code in ((path_a, "Akash", "C-11"), (path_t, "Tharindu", "C-22")):
                ws = openpyxl.load_workbook(path).active
                self.assertEqual({ws.cell(r, 2).value for r in range(8, ws.max_row + 1)}, {user})
                self.assertEqual({ws.cell(r, 4).value for r in range(8, ws.max_row + 1)}, {code})
            ok, _, _, _ = export("Tharindu", "C-22", [header1])
            self.assertTrue(ok)
            ws = openpyxl.load_workbook(path_t).active
            self.assertEqual(ws.max_row, 15)


if __name__ == "__main__":
    unittest.main()
